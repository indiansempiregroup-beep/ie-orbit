#!/usr/bin/env python3
"""Generate ops-mobile Add/Edit form inventory workbook. Customer-mobile excluded."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference

OUT = Path(__file__).with_name("ops-mobile-add-edit-forms.xlsx")

# area, screen, route, mode, how_to_open, section, field, key, control, required, validation, error, conditional
ROWS = [
    # Auth / onboarding
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "First name", "firstName", "text", "Yes", "trim required", "First name is required", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "Last name", "lastName", "text", "Yes", "trim required", "Last name is required", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "Email", "email", "email", "Yes", "required + email pattern", "Email is required / Invalid email address: {value}", "Locked if Google sign-up"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "Mobile", "mobile", "phone", "Yes", "Indian mobile (+91)?[6-9]XXXXXXXXX", "Phone number is required / Enter a valid 10-digit Indian mobile number", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "Owner email OTP", "ownerOtpCode", "number", "Conditional", "6-digit after send code", "Sign-in code is required / invalid", "Hidden when Google token present; send code to owner email"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 0 · You", "Affiliate code", "affiliateCode", "text", "No", "none", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Business name", "businessName", "text", "Yes", "trim required", "Business name is required", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Display name", "displayName", "text", "Yes", "trim required", "Display name is required", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Business email", "businessEmail", "email", "Yes", "required + email pattern", "Email is required / Invalid email address: {value}", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Phone", "businessPhone", "phone", "Yes", "Indian mobile", "Phone number is required / invalid Indian mobile", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Address", "address + lat/lng", "places + map", "Yes", "address + coordinates required", "Address is required; map pin required in later checks", "City/state/country/postal locked when coords set"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "City", "city", "text", "Yes", "trim required", "City is required", "Locked with coords"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "State", "state", "text", "No", "none", "", "Locked with coords"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Country code", "country", "text", "Yes", "trim required", "Country is required", "Locked with coords"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 1 · Business", "Postal code", "postalCode", "text", "Yes", "trim required", "Postal code is required", "Locked with coords"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Timezone", "timezone", "select", "Yes (UI)", "always has default", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Currency", "currency", "select", "Yes (UI)", "always has default", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Language", "language", "select", "Yes (UI)", "always has default", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Products", "selectedProducts", "multi chip", "Yes", "at least 1 product", "Select at least one product (step gate)", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Package per product", "planCodes[product]", "cards", "Yes", "plan required per selected product", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Skip hours", "skipHours", "switch", "No", "none", "", "If off: ≥1 open day; close > open"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 2 · Setup", "Day open + Opens/Closes", "businessHours[day]", "switch + time", "Conditional", "close after open when day is open", "", "Hidden when skip hours; times only if day open"),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 3 · Brand", "Primary color", "primaryColor", "text/color", "No", "none", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 3 · Brand", "Secondary color", "secondaryColor", "text/color", "No", "none", "", ""),
    ("Auth", "Register business", "RegisterWizard", "Add", "Login → Create workspace", "Step 3 · Brand", "Logo", "logoAsset", "image", "No", "none", "", ""),
    ("Auth", "Accept invitation", "AcceptInvitation", "Add", "Invite email link / token", "Join", "First name", "firstName", "text", "Yes", "trim required", "First name is required", "No form if token missing"),
    ("Auth", "Accept invitation", "AcceptInvitation", "Add", "Invite email link / token", "Join", "Last name", "lastName", "text", "No", "none", "", "No password — OTP sign-in after accept"),
    ("Auth", "Sign in", "Login", "Update", "Login", "OTP", "Email", "email", "email", "Yes", "required + email pattern", "Email is required / Invalid email address: {value}", "Sign in with OTP"),
    ("Auth", "Sign in", "Login", "Update", "Login", "OTP", "Sign-in code", "code", "number", "Yes", "6-digit", "Sign-in code is required", "After send code"),
    ("Auth", "Sign in", "Login", "Update", "Login", "Mobile OTP", "Mobile number", "phone", "phone", "Conditional", "when WhatsApp enabled", "Mobile number is required", "Platform WhatsApp sender configured"),

    # Core entity
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Identity", "Display name", "displayName", "text", "Conditional", "display name OR first+last", "Display name or first name is required", "Edit if customerId"),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Identity", "First name", "firstName", "text", "No", "counts toward identity with last name", "", ""),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Identity", "Last name", "lastName", "text", "No", "none", "", ""),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Contact", "Email", "email", "email", "Conditional", "optional format; email OR phone required", "Invalid email address: {value} / Email or phone is required", ""),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Contact", "Phone", "phone", "phone", "Conditional", "optional Indian mobile; email OR phone required", "Enter a valid 10-digit Indian mobile number / Email or phone is required", ""),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "GST (Orbit Mart)", "GSTIN", "gstin", "text (max 15)", "No", "empty OK else 15-char format + checksum", "GSTIN must be 15 characters / Invalid format / Invalid check digit", "Only if Shopie / Orbit Mart subscribed"),
    ("Customers", "Add/Edit customer", "CustomerForm", "Add + Edit", "More → Customers → + or Edit", "Address", "Search address", "address.line1 + lat/lng", "places + map", "No", "none", "", "Optional"),
    ("Customers", "Customer repay (nested)", "CustomerDetail", "Add", "Customers → open customer with balance", "Borrow repayment", "Amount", "amount", "number", "Yes (UI)", "used when balance > 0", "", "Only if outstanding balance"),
    ("Customers", "Customer repay (nested)", "CustomerDetail", "Add", "Customers → open customer with balance", "Borrow repayment", "Method", "method", "chips cash|upi|card", "No", "none", "", ""),
    ("Customers", "Customer repay (nested)", "CustomerDetail", "Add", "Customers → open customer with balance", "Borrow repayment", "Note", "note", "text", "No", "none", "", ""),

    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Profile", "Profile photo", "photoAsset", "image", "No", "none", "", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Profile", "First name", "firstName", "text", "Conditional", "required unless display name filled", "First name is required", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Profile", "Last name", "lastName", "text", "No", "none", "", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Profile", "Display name", "displayName", "text", "No", "can substitute for first name", "", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Contact", "Email", "email", "email", "Conditional", "required if Send invite; else optional format", "Email is required / Invalid email address: {value}", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Contact", "Phone", "phone", "phone", "No", "optional Indian mobile", "Enter a valid 10-digit Indian mobile number", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Access", "Available for bookings", "isBookable", "checkbox", "No", "none", "", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Access", "Role", "role", "cards staff|manager", "No", "defaulted", "", ""),
    ("Staff", "Add/Edit staff", "StaffForm", "Add + Edit", "More → Staff → + or Edit", "Access", "Send login invitation", "sendInvite", "checkbox", "No", "none", "", "Hidden on edit if user already linked"),
    ("Staff", "Weekly schedule", "StaffSchedule", "Edit", "Staff detail → Weekly schedule", "Mon–Sun", "Available", "days[].available", "switch", "No", "no close>open check", "", ""),
    ("Staff", "Weekly schedule", "StaffSchedule", "Edit", "Staff detail → Weekly schedule", "Mon–Sun", "Start / End", "days[].start/end", "time", "Conditional", "shown when available", "", "Only if day available"),
    ("Staff", "Staff availability · Leave", "StaffAvailability", "Add", "Staff detail → Availability", "Leave", "Days", "leaveDays", "multi calendar", "Yes", "≥1 day", "", ""),
    ("Staff", "Staff availability · Leave", "StaffAvailability", "Add", "Staff detail → Availability", "Leave", "Half / Full day", "leaveKind", "chips", "No", "none", "", ""),
    ("Staff", "Staff availability · Leave", "StaffAvailability", "Add", "Staff detail → Availability", "Leave", "Reason", "leaveReason", "text", "No", "none", "", ""),
    ("Staff", "Staff availability · Extra", "StaffAvailability", "Add", "Staff detail → Availability", "Extra hours", "Starts", "specialStart", "datetime", "Yes (UI)", "end must be after start", "", ""),
    ("Staff", "Staff availability · Extra", "StaffAvailability", "Add", "Staff detail → Availability", "Extra hours", "Ends", "specialEnd", "datetime", "Yes (UI)", "end > start", "", ""),
    ("Staff", "Staff availability · Block", "StaffAvailability", "Add", "Staff detail → Availability", "Block slot", "Date", "slotDate", "date", "Yes (UI)", "none extra", "", ""),
    ("Staff", "Staff availability · Block", "StaffAvailability", "Add", "Staff detail → Availability", "Block slot", "Start / End", "slotStart/slotEnd", "time", "Yes (UI)", "none extra", "", ""),
    ("Staff", "Staff availability · Block", "StaffAvailability", "Add", "Staff detail → Availability", "Block slot", "Reason", "slotReason", "text", "No", "none", "", ""),
    ("Staff", "Staff availability · Emergency", "StaffAvailability", "Add", "Staff detail → Availability", "Emergency open", "Date + Start/End + Reason", "slotDate/start/end/reason", "date/time/text", "Date/times Yes (UI)", "same as block", "", ""),
    ("Staff", "Staff availability · Services", "StaffAvailability", "Add", "Staff detail → Availability", "Services", "Service", "serviceId", "select", "Yes", "required to assign", "", ""),

    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Basics", "Service image", "imageAsset", "image", "No", "none", "", ""),
    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Basics", "Name", "name", "text", "Yes", "trim required", "Name is required", ""),
    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Basics", "Description", "description", "multiline", "No", "none", "", ""),
    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Duration & price", "Duration", "duration", "select 15–240 min", "Yes", "must be > 0", "Duration is required", ""),
    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Duration & price", "Price", "price", "decimal", "No", "optional; sent if filled", "", ""),
    ("Services", "Add/Edit service", "ServiceForm", "Add + Edit", "More → Services → + or Edit", "Duration & price", "Points earned on complete", "loyaltyPointsEarn", "number", "No", "coerced ≥ 0", "", ""),

    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "Office name", "name", "text", "Yes", "trim required", "Office name is required.", ""),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "Address", "address + lat/lng", "places + map", "Yes", "full Places result + map location", "Select a full office address from Google Places. / Google Map location is required.", ""),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "City", "city", "text", "Yes", "required", "", "Locked with coords"),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "State", "state", "text", "No", "none", "", "Locked with coords"),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "Country", "country", "text", "Yes", "required", "", "Locked with coords"),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "Postal code", "postalCode", "text", "No", "none", "", "Locked with coords"),
    ("Offices", "Add/Edit office", "BranchForm", "Add + Edit", "Settings → Offices → + or Edit", "Office", "Phone (rider contact)", "phone", "phone", "No", "none", "", ""),
    ("Offices", "Add/Edit office", "BranchForm", "Edit", "Settings → Offices → Edit", "Actions", "Set primary / Deactivate", "—", "buttons", "—", "cannot deactivate last active office", "", "Edit only"),

    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Owner", "Customer", "customerId", "select", "Yes", "required", "owner required message", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Pet photo", "photoAsset", "image", "No", "none", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Pet name", "name", "text", "Yes", "trim required", "name required message", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Species", "species", "chips", "Yes (UI)", "default Dog", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Breed", "breed", "text", "No", "none", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Sex", "sex", "chips", "No", "none", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Birthday", "birthday", "date", "No", "YYYY-MM-DD if set", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Add + Edit", "More → Pets → + or Edit", "Pet", "Care notes", "notes", "multiline", "No", "none", "", ""),
    ("Pets", "Add/Edit pet", "ShopPetForm", "Edit", "Pets → Edit", "Actions", "Delete pet", "—", "button", "—", "none", "", "Edit only"),
    ("Pets", "Notify owner (nested)", "ShopPetDetail", "Add", "Pets → pet → Notify", "Message", "Subject", "subject", "text", "Yes", "both required", "Subject and message are required.", ""),
    ("Pets", "Notify owner (nested)", "ShopPetDetail", "Add", "Pets → pet → Notify", "Message", "Message", "body", "multiline", "Yes", "both required", "Subject and message are required.", ""),

    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Lookup", "Search by name", "nameLookup", "text + search", "No", "none", "", "Prefill only"),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Lookup", "Scan barcode", "barcode", "scanner", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Photos", "Photos", "images", "image slots", "No", "JS type checks in picker", "", "Up to N images"),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "Product name", "name", "text", "Yes", "trim required", "Product name is required", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "Brand", "brand", "text", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "SKU", "sku", "text", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "Pack size / quantity", "pack_size", "text", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "Status", "status", "select", "Yes", "trim required", "Status is required", "Default active"),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Product", "Category", "category", "chips", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Price & tax", "Price", "price", "decimal", "Yes", "required; finite number ≥ 0", "Price is required / Enter a valid price", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Price & tax", "GST %", "tax_rate", "decimal", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Price & tax", "Currency", "currency", "select", "Yes", "trim required", "Currency is required", "Default INR"),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Price & tax", "GST included / excluded", "tax_inclusive", "chips", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Stock", "Stock on hand", "stock_on_hand", "number", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Stock", "Low stock", "low_stock_threshold", "number", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Stock", "Godown", "godownId", "select", "Conditional", "none extra", "", "If godowns plan feature"),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Barcode", "Barcode", "barcode", "text", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Barcode", "Barcode type", "barcode_type", "chips", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Details", "Description", "description", "multiline", "No", "none", "", ""),
    ("Shop catalog", "Add/Edit product", "ShopProductAdd", "Add + Edit", "More → Products → + or product", "Details", "HTML details", "details_html", "html editor", "No", "none", "", ""),
    ("Shop catalog", "Add many products", "ShopProductsAddMany", "Add", "Products → Add many", "Defaults", "GST %, HSN, category, status, godown", "defaults.*", "mixed", "No", "applied to rows", "", ""),
    ("Shop catalog", "Add many products", "ShopProductsAddMany", "Add", "Products → Add many", "Row", "Name", "rows[].name", "text", "Yes if row has content", "required; duplicate in list; already in catalog", "Name is required. / Duplicate in this list. / Already in catalog: …", ""),
    ("Shop catalog", "Add many products", "ShopProductsAddMany", "Add", "Products → Add many", "Row", "Barcode / Price / GST / Stock / SKU / Brand / Pack / HSN / Category / Status", "rows[].*", "mixed", "No", "none extra", "", ""),

    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "Who", "Customer", "customerId", "select", "Yes", "required", "Customer is required", ""),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "What", "Services", "selectedServiceIds", "multi", "Yes", "≥1", "Select at least one service.", ""),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "Where", "Office", "branchId", "select", "Conditional", "required if more than one office", "Select an office…", "Hidden if only one branch"),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "Who", "Staff", "staffId", "select", "No", "none", "", "Hidden for staff role / cleared if multi-specialist"),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "When", "Date", "date", "chips + calendar", "Yes (UI)", "none extra", "", ""),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "When", "Time slot", "selectedSlot", "slot grid", "Yes", "required", "Select an available time slot.", ""),
    ("Bookings", "New booking", "CreateBooking", "Add", "Bookings → + / New booking", "Notes", "Notes", "notes", "multiline", "No", "none", "", ""),
    ("Bookings", "Booking detail actions", "BookingDetail", "Edit", "Open a booking", "Cancel / etc.", "Reason", "reason", "text", "No", "none", "", ""),
    ("Bookings", "Booking detail actions", "BookingDetail", "Edit", "Open a booking", "Reschedule", "Date + slot", "date + slot", "calendar + slots", "Slot yes", "slot required to confirm", "", ""),
    ("Bookings", "Booking detail actions", "BookingDetail", "Edit", "Open a booking", "Reassign", "Staff", "staffId(s)", "select", "No", "per-line if multi-service", "", ""),

    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "Photo", "photoAsset", "image", "No", "none", "", ""),
    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "First name", "firstName", "text", "Yes", "trim required", "First name is required", ""),
    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "Last name", "lastName", "text", "No", "none", "", ""),
    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "Phone", "phone", "phone", "No", "optional Indian mobile", "Enter a valid 10-digit Indian mobile number", ""),
    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "Language", "language", "select", "Yes (UI)", "always valued", "", ""),
    ("Profile", "Edit profile", "ProfileEdit", "Edit", "More → Profile → Edit", "You", "Timezone", "timezone", "select", "Yes (UI)", "always valued", "", ""),
    ("Profile", "Security", "Profile", "Edit", "Profile → Security section", "Account", "Biometric login", "biometricEnabled", "switch", "No", "device capability", "", "OTP-only — no password fields"),
    ("Profile", "Notification preferences", "NotificationPreferences", "Edit", "Profile → Notifications", "Channels", "Email / Push", "email, push", "switches", "No", "none", "", ""),

    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Branding", "Business logo", "logoAsset", "image", "No", "none", "", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Branding", "Legal / business name", "businessName", "text", "Conditional", "legal OR display required", "Business name is required", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Branding", "Display name", "displayName", "text", "Conditional", "same identity rule", "", "UI marked required"),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Contact", "Business email", "email", "email", "No", "optional format", "Invalid email address: {value}", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Contact", "Primary contact", "primaryContact", "phone", "No", "optional Indian mobile", "Enter a valid 10-digit Indian mobile number", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Contact", "Website", "website", "url", "No", "optional URL", "Enter a valid website URL (for example, https://yoursalon.com)", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Location", "Address + city/state/postal/country", "address_*", "places + text", "No", "none", "", "Derived locked with coords"),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Tax & GST", "GSTIN", "gst_tax_number", "text (max 15)", "No", "empty OK else format + checksum", "GSTIN must be 15 characters / Invalid format / Invalid check digit", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Regional", "Timezone", "timezone", "select", "Yes (UI)", "always valued", "", ""),
    ("Settings", "Edit business", "BusinessEdit", "Edit", "Settings → Business profile → Edit", "Regional", "Currency", "currency", "select", "Yes (UI)", "always valued", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Razorpay", "Enabled", "enabled", "switch", "No", "none", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Razorpay", "Key ID", "keyId", "text", "Yes (UI)", "trim only", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Razorpay", "Key Secret", "keySecret", "password", "Conditional", "required if not already stored", "", "Blank keeps saved secret"),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Razorpay", "Webhook Secret", "webhookSecret", "password", "No", "none", "", "Blank keeps saved"),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Cashfree", "App ID / Secret / enabled", "cashfree.*", "text/password/switch", "Secret if not stored", "trim only", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Other", "COD enabled", "codEnabled", "switch", "No", "none", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Other", "UPI ID", "upiVpa", "text", "No", "none", "", ""),
    ("Settings", "Payments", "PaymentSettings", "Edit", "Settings → Payments", "Other", "Static payment QR", "qr", "image", "No", "none", "", ""),
    ("Settings", "Products & extras", "ProductSettings", "Edit", "Settings → Products & billing", "Loyalty", "Enable + points per ₹1 / max % / min redeem / points per ₹100", "loyalty.*", "switch + numbers", "No", "coerced on save (points ≥1, max % 0–100)", "", "Plan gated"),
    ("Settings", "Products & extras", "ProductSettings", "Edit", "Settings → Products & billing", "Add-ons", "Extra staff / offices", "extras", "number", "No", "none", "", ""),
    ("Settings", "Team invite", "Team", "Add", "Settings → Team", "Invite", "Email", "email", "email", "Yes", "required + pattern", "Email is required / Invalid email address: {value}", ""),
    ("Settings", "Team invite", "Team", "Add", "Settings → Team", "Invite", "Role", "role", "chips staff|manager", "No", "none", "", ""),

    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Code", "code", "text", "Yes", "trim required", "Code is required", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Name", "name", "text", "Yes", "trim required", "Name is required", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Description", "description", "text", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Percent | Amount", "discountType", "chips", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Percent / Amount value", "discountValue", "number", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Min. order", "minOrder", "number", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Max discount", "maxDiscount", "number", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Starts / Ends", "startsAt/endsAt", "date", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "Total uses / Uses per customer", "totalUses/perCustomer", "number", "No", "none", "", ""),
    ("Shop ops", "Coupons", "ShopCoupons", "Add + Edit", "More → Coupons → Add/edit", "Coupon", "First order only / Active", "firstOrder/active", "switches", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Zone name", "name", "text", "Yes", "trim required", "Zone name is required", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Cities", "cities", "CSV text", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Postal prefixes", "postalPrefixes", "CSV text", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Delivery fee", "fee", "number", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Minimum order", "minOrder", "number", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Notes", "notes", "text", "No", "none", "", ""),
    ("Shop ops", "Delivery zones", "ShopDeliveryZones", "Add + Edit", "More → Delivery zones", "Zone", "Same-day / Deliver now / Enabled", "flags", "switches", "No", "none", "", ""),
    ("Shop ops", "Delivery settings", "ShopDeliverySettings", "Edit", "More → Delivery settings", "Provider", "Instant delivery", "instant", "switch", "No", "none", "", ""),
    ("Shop ops", "Delivery settings", "ShopDeliverySettings", "Edit", "More → Delivery settings", "Provider", "Provider", "provider", "select", "No", "none", "", ""),
    ("Shop ops", "Delivery settings", "ShopDeliverySettings", "Edit", "More → Delivery settings", "Credentials", "API base / Porter key / Shiprocket email+password+pickup+weight / webhook", "creds.*", "text/password", "Porter key if not stored", "trim; credentials when provider ≠ mock", "", "Shiprocket block if shiprocket_quick; absorb cap if split"),
    ("Shop ops", "Delivery settings", "ShopDeliverySettings", "Edit", "More → Delivery settings", "Fees", "Who pays fee / Free above / Absorb cap", "chargeBearer etc", "select/number", "No", "absorb cap if split", "", ""),
    ("Shop ops", "Stock adjust", "ShopStockAdjust", "Edit", "Products / stock → Adjust", "Adjust", "Office / Godown", "office/godown", "select", "Conditional", "plan/offices", "", ""),
    ("Shop ops", "Stock adjust", "ShopStockAdjust", "Edit", "Products / stock → Adjust", "Adjust", "Quantity change (+/-)", "qty", "number", "Yes", "non-zero", "Enter a non-zero quantity change", ""),
    ("Shop ops", "Stock adjust", "ShopStockAdjust", "Edit", "Products / stock → Adjust", "Adjust", "Movement type", "type", "select", "No", "none", "", ""),
    ("Shop ops", "Stock adjust", "ShopStockAdjust", "Edit", "Products / stock → Adjust", "Adjust", "Reason", "reason", "text", "No", "none", "", ""),
    ("Shop ops", "Godowns", "ShopGodowns", "Add + Edit", "Books / settings → Godowns", "Godown", "Name", "name", "text", "Yes", "required", "", ""),
    ("Shop ops", "Godowns", "ShopGodowns", "Add + Edit", "Books / settings → Godowns", "Godown", "Address", "address + pin", "places", "Yes", "required + pin", "", ""),
    ("Shop ops", "Godowns", "ShopGodowns", "Add + Edit", "Books / settings → Godowns", "Godown", "Pickup phone", "phone", "phone", "No", "none", "", ""),
    ("Shop ops", "Godowns", "ShopGodowns", "Add + Edit", "Books / settings → Godowns", "Godown", "Code", "code", "text", "No", "none", "", ""),
    ("Shop ops", "Godowns", "ShopGodowns", "Add + Edit", "Books / settings → Godowns", "Godown", "Default", "isDefault", "switch", "No", "none", "", ""),
    ("Shop ops", "Godown transfer", "ShopGodowns", "Add", "Godowns → Transfer", "Transfer", "From / To / Product / Quantity", "from/to/product/qty", "selects + number", "Yes", "all required; from ≠ to", "", ""),
    ("Shop ops", "Reward points", "ShopLoyalty", "Edit", "More / Shop → Reward points", "Program", "Enable", "enabled", "switch", "No", "none", "", ""),
    ("Shop ops", "Reward points", "ShopLoyalty", "Edit", "More / Shop → Reward points", "Program", "Points per ₹1 / Max redeem % / Min redeem / Points per ₹100", "metrics", "digits", "No", "coerced on save", "", ""),
    ("Shop ops", "Order tracking (nested)", "ShopOrderDetail", "Edit", "Orders → order", "Tracking", "Carrier", "carrier", "text", "No", "none", "", ""),
    ("Shop ops", "Order tracking (nested)", "ShopOrderDetail", "Edit", "Orders → order", "Tracking", "AWB", "awb", "text", "Yes", "required", "AWB or tracking number is required.", ""),
    ("Shop ops", "Order tracking (nested)", "ShopOrderDetail", "Edit", "Orders → order", "Tracking", "ETA / Notify", "eta/notify", "date / switch", "No", "none", "", ""),
    ("Shop ops", "Order return (nested)", "ShopOrderDetail", "Add", "Orders → order → Return", "Return", "Per-line qty", "qtyByLine", "stepper", "Yes", "≥1 qty", "", ""),
    ("Shop ops", "Order return (nested)", "ShopOrderDetail", "Add", "Orders → order → Return", "Return", "Reason", "reason", "text", "No", "none", "", ""),

    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Supplier name", "name", "text", "Yes", "trim required", "Supplier name is required", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Phone", "phone", "phone", "No", "optional Indian mobile", "Enter a valid 10-digit Indian mobile number", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Email", "email", "email", "No", "optional format", "Invalid email address: {value}", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "GSTIN", "gstin", "text (max 15)", "No", "empty OK else checksum", "GSTIN format / check digit messages", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Billing state", "billingState", "text", "No", "none", "", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Billing address", "billingAddress", "text", "No", "none", "", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Credit limit", "creditLimit", "number", "No", "none", "", ""),
    ("Books", "Add supplier", "ShopBooksParties", "Add", "Books → Parties → +", "Supplier", "Opening balance", "openingBalance", "number", "No", "none", "", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Kind", "kind", "chips expense|income", "No", "none", "", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Category", "category", "select", "Yes", "category + amount > 0", "Enter a category and amount", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Category name", "categoryName", "text", "No", "none", "", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Amount", "amount", "number", "Yes", "> 0", "Enter a category and amount", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Paid from / Received into", "account", "select", "No", "none", "", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Date", "date", "date", "Yes (UI)", "none", "", ""),
    ("Books", "Expense / other income", "ShopBooksExpense", "Add", "Books → Expense", "Entry", "Notes", "notes", "text", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Account", "Name", "accountName", "text", "Yes", "required", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Account", "Cash | Bank", "kind", "chips", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Account", "Opening balance", "opening", "number", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Payment", "In | Out", "direction", "chips", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Payment", "Customer | Supplier", "party", "select", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Payment", "Account", "account", "select", "No", "none", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Payment", "Amount", "amount", "number", "Yes", "> 0", "", ""),
    ("Books", "Cash / bank", "ShopBooksCash", "Add", "Books → Cash & bank", "Payment", "Date / Notes", "date/notes", "date/text", "Date Yes (UI)", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Direction in|out", "direction", "chips", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Customer | Supplier", "party", "select", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Cheque number", "number", "text", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Amount", "amount", "number", "Yes", "> 0", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Bank", "bank", "text", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Due date", "dueDate", "date", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Cash/bank account", "account", "select", "No", "none", "", ""),
    ("Books", "Cheques", "ShopBooksCheques", "Add", "Books → Cheques", "Cheque", "Notes", "notes", "text", "No", "none", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Loan", "Customer", "customer", "select", "No", "none", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Loan", "Title", "title", "text", "No", "none", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Loan", "Principal", "principal", "number", "Yes", "> 0", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Loan", "Interest %", "interest", "number", "No", "none", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Loan", "Start date / Notes", "start/notes", "date/text", "No", "none", "", ""),
    ("Books", "Loans", "ShopBooksLoans", "Add", "Books → Loans", "Repay", "Repay amount", "repayAmount", "number", "Yes if repay", "> 0", "", "Inline on existing loan"),
    ("Books", "Documents", "ShopBooksDocuments", "Add", "Books → Documents (type in params)", "Header", "Customer | Supplier", "party", "select", "No", "none", "", "Depends on doc type"),
    ("Books", "Documents", "ShopBooksDocuments", "Add", "Books → Documents (type in params)", "Header", "Date", "date", "date", "Yes (UI)", "none", "", ""),
    ("Books", "Documents", "ShopBooksDocuments", "Add", "Books → Documents (type in params)", "Lines", "Product / Qty / Rate / GST %", "lines[]", "mixed", "Yes", "≥1 valid line", "", ""),
    ("Books", "Documents", "ShopBooksDocuments", "Add", "Books → Documents (type in params)", "Header", "Notes", "notes", "text", "No", "none", "", ""),
    ("Books", "GST compliance", "ShopBooksCompliance", "Edit", "Books → Compliance", "Toggles", "E-invoice / E-way", "switches", "switch", "No", "none", "", ""),
    ("Books", "GST compliance", "ShopBooksCompliance", "Edit", "Books → Compliance", "Provider", "Provider + credentials", "provider/user/pass", "select/text", "Creds if not mock", "username/password required UI if not mock", "", ""),
    ("Books", "GST compliance", "ShopBooksCompliance", "Edit", "Books → Compliance", "Seller", "Legal name, addr1, city, PIN, state code", "seller.*", "text", "Yes", "required", "", ""),
    ("Books", "GST compliance", "ShopBooksCompliance", "Edit", "Books → Compliance", "Seller", "Trade name / addr2", "seller.optional", "text", "No", "none", "", ""),
    ("Books", "Sale e-way / cancel (nested)", "ShopBooksSale", "Add", "Books → Sale voucher", "E-way", "Vehicle number", "vehicle", "text", "Yes", "required", "", ""),
    ("Books", "Sale e-way / cancel (nested)", "ShopBooksSale", "Add", "Books → Sale voucher", "E-way", "Transport mode / Distance / Transporter", "mode/distance/name", "mixed", "No", "none", "", ""),
    ("Books", "Sale e-way / cancel (nested)", "ShopBooksSale", "Add", "Books → Sale voucher", "Cancel", "Cancel reason", "reason", "text", "Yes", "required for e-invoice/e-way cancel", "", ""),

    ("Grow", "WhatsApp", "GrowWhatsApp", "Edit + send", "More → Grow → WhatsApp", "Send", "Send-to mode", "mode", "chips", "No", "none", "", ""),
    ("Grow", "WhatsApp", "GrowWhatsApp", "Edit + send", "More → Grow → WhatsApp", "Send", "Customer / Supplier", "party", "select", "Conditional", "if not manual", "", ""),
    ("Grow", "WhatsApp", "GrowWhatsApp", "Edit + send", "More → Grow → WhatsApp", "Send", "Country dial + Phone", "phone", "select + phone", "Yes to open WA", "national number required to open WhatsApp", "", ""),
    ("Grow", "WhatsApp", "GrowWhatsApp", "Edit + send", "More → Grow → WhatsApp", "Send", "Default message", "message", "multiline", "No", "none", "", ""),
    ("Grow", "WhatsApp", "GrowWhatsApp", "Edit + send", "More → Grow → WhatsApp", "Send", "Attachment image", "image", "image", "No", "none", "", ""),
    ("Grow", "Google Profile", "GrowGoogleProfile", "Edit", "More → Grow → Google Profile", "Place", "Profile URL", "url", "url", "Conditional", "open requires URL or Place ID", "", ""),
    ("Grow", "Google Profile", "GrowGoogleProfile", "Edit", "More → Grow → Google Profile", "Place", "Place ID", "placeId", "text", "Conditional", "one of URL / Place ID", "", ""),
    ("Grow", "Ads", "GrowAds", "Add + Edit", "More → Grow → Ads", "Ad", "Title", "title", "text", "Yes", "trim required", "", "Max 5 active ads"),
    ("Grow", "Ads", "GrowAds", "Add + Edit", "More → Grow → Ads", "Ad", "Short message", "message", "text", "No", "none", "", ""),
    ("Grow", "Ads", "GrowAds", "Add + Edit", "More → Grow → Ads", "Ad", "Link", "link", "url", "No", "none", "", ""),
    ("Grow", "Ads", "GrowAds", "Add + Edit", "More → Grow → Ads", "Ad", "Active / Inactive", "isActive", "chips", "No", "none", "", ""),
    ("Grow", "Ads", "GrowAds", "Add + Edit", "More → Grow → Ads", "Ad", "Ad image", "image", "image", "No", "none", "", ""),
    ("Grow", "Referrals", "GrowReferral", "Edit", "More → Grow → Referrals", "Program", "Enabled / Disabled", "enabled", "chips", "No", "none", "", ""),
    ("Grow", "Referrals", "GrowReferral", "Edit", "More → Grow → Referrals", "Program", "Points", "points", "number", "Yes", "≥ 0 numeric", "", ""),
    ("Grow", "Referrals", "GrowReferral", "Edit", "More → Grow → Referrals", "Program", "Success event", "event", "chips", "No", "none", "", ""),

    ("Admin", "Suspend / reactivate tenant", "PlatformAdminTenantDetail", "Edit", "Platform Admin → Tenant", "Action", "Reason", "reason", "text", "Yes", "required", "", "Platform admin only"),
    ("Support", "Ticket reply", "SupportTicketDetail", "Add", "Support tickets → ticket", "Reply", "Reply / Internal note", "body", "multiline", "Yes", "send disabled if empty", "", "Internal toggle platform-only"),
]

HEADERS = [
    "Area",
    "Screen",
    "Route",
    "Mode",
    "How to open in ops-mobile",
    "Section",
    "Field label",
    "Form key",
    "Control",
    "Required",
    "Client validation (current)",
    "Error message",
    "Conditional / notes",
    "Your change (keep / edit / drop)",
    "Updated required?",
    "Updated validation / message",
]

HELP = [
    ["Ops-mobile Add/Edit form inventory"],
    ["Generated from current ops-mobile source (apps/ops-mobile). Customer-mobile is excluded."],
    ["Expo: Metro is on exp://LAN:8082 and web http://localhost:8082. Authenticated screens need a signed-in session; this workbook is the field/validation catalog to review."],
    [""],
    ["How to review"],
    ["1. Open the Fields sheet. Filter by Area or Screen."],
    ["2. Fill Your change: keep / edit / drop (dropdown)."],
    ["3. If edit, fill Updated required? and Updated validation / message."],
    ["4. Send the filled file back so the forms can be updated."],
    [""],
    ["Shared validation helpers"],
    ["Email", "Empty (if required) → Email is required. Invalid → Invalid email address: {typed value}. Pattern: not empty, has @ and a dot-domain."],
    ["Indian mobile", "Optional unless required. Compact form (+91)? then 10 digits starting 6–9. Message: Enter a valid 10-digit Indian mobile number."],
    ["Sign-in OTP", "6-digit email (or WhatsApp) code for auth. Not used for Razorpay/Shiprocket integration secrets."],
    ["Website (ops)", "Optional. Must parse as URL (https:// added if missing)."],
    ["GSTIN (Shopie)", "Empty allowed. Else 15 chars, state+PAN pattern, checksum."],
    [""],
    ["Sheets"],
    ["Fields", "One row per field. This is the review sheet."],
    ["Screens", "One row per Add/Edit screen / nested form."],
    ["By area", "Count of fields per area (chart)."],
]


def style_header(ws, n):
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    for col in range(1, n + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = thin
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    wb = Workbook()

    help_ws = wb.active
    help_ws.title = "How to use"
    help_ws["A1"] = HELP[0][0]
    help_ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    for i, row in enumerate(HELP[1:], 2):
        for j, val in enumerate(row, 1):
            help_ws.cell(i, j, val)
            if j == 1 and val in {
                "How to review",
                "Shared validation helpers",
                "Sheets",
                "Email",
                "Indian mobile",
                "Password",
                "Website (ops)",
                "GSTIN (Shopie)",
                "Fields",
                "Screens",
                "By area",
            }:
                help_ws.cell(i, j).font = Font(bold=True)
    help_ws.column_dimensions["A"].width = 28
    help_ws.column_dimensions["B"].width = 110
    help_ws.row_dimensions[1].height = 24

    fields = wb.create_sheet("Fields")
    fields.append(HEADERS)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    alt = PatternFill("solid", fgColor="F8FAFC")
    review_fill = PatternFill("solid", fgColor="FFF7ED")
    for i, row in enumerate(ROWS, 2):
        for j, val in enumerate(row, 1):
            cell = fields.cell(i, j, val)
            cell.alignment = wrap
            cell.border = thin
            if i % 2 == 0:
                cell.fill = alt
        for j in range(14, 17):
            fields.cell(i, j).fill = review_fill
            fields.cell(i, j).border = thin
        fields.row_dimensions[i].height = 36

    style_header(fields, len(HEADERS))
    autosize(fields, [14, 28, 26, 14, 36, 22, 36, 24, 18, 14, 42, 42, 36, 22, 18, 36])
    last = fields.max_row
    dv = DataValidation(type="list", formula1='"keep,edit,drop"', allow_blank=True)
    dv.error = "Use keep, edit, or drop"
    dv.errorTitle = "Invalid"
    dv.prompt = "keep / edit / drop"
    dv.promptTitle = "Your change"
    fields.add_data_validation(dv)
    dv.add(f"N2:N{last}")
    dv2 = DataValidation(type="list", formula1='"Yes,No,Conditional"', allow_blank=True)
    fields.add_data_validation(dv2)
    dv2.add(f"O2:O{last}")

    # Screens unique
    screens = wb.create_sheet("Screens")
    sh_headers = ["Area", "Screen", "Route", "Mode", "How to open", "Field count", "File hint"]
    screens.append(sh_headers)
    seen = {}
    for r in ROWS:
        key = (r[0], r[1], r[2], r[3], r[4])
        seen[key] = seen.get(key, 0) + 1
    file_hint = {
        "RegisterWizard": "features/onboarding/RegisterWizardScreen.tsx",
        "AcceptInvitation": "features/auth/AcceptInvitationScreen.tsx",
        "Login": "features/auth/LoginScreen.tsx",
        "CustomerForm": "features/customers/CustomerFormScreen.tsx",
        "CustomerDetail": "features/customers/CustomerDetailScreen.tsx",
        "StaffForm": "features/staff/StaffFormScreen.tsx",
        "StaffSchedule": "features/staff/StaffScheduleScreen.tsx",
        "StaffAvailability": "features/staff/StaffAvailabilityScreen.tsx",
        "ServiceForm": "features/services/ServiceFormScreen.tsx",
        "BranchForm": "features/branches/BranchFormScreen.tsx",
        "ShopPetForm": "features/shop/ShopPetFormScreen.tsx",
        "ShopPetDetail": "features/shop/ShopPetDetailScreen.tsx",
        "ShopProductAdd": "features/shop/ShopProductAddScreen.tsx",
        "ShopProductsAddMany": "features/shop/ShopProductsAddManyScreen.tsx",
        "CreateBooking": "features/bookings/CreateBookingScreen.tsx",
        "BookingDetail": "features/bookings/BookingDetailScreen.tsx",
        "ProfileEdit": "features/profile/ProfileEditScreen.tsx",
        "NotificationPreferences": "features/profile/NotificationPreferencesScreen.tsx",
        "BusinessEdit": "features/settings/BusinessEditScreen.tsx",
        "PaymentSettings": "features/settings/PaymentSettingsScreen.tsx",
        "ProductSettings": "features/settings/ProductSettingsScreen.tsx",
        "Team": "features/team/TeamScreen.tsx",
        "ShopCoupons": "features/shop/ShopCouponsScreen.tsx",
        "ShopDeliveryZones": "features/shop/ShopDeliveryZonesScreen.tsx",
        "ShopDeliverySettings": "features/shop/ShopDeliverySettingsScreen.tsx",
        "ShopStockAdjust": "features/shop/ShopStockAdjustScreen.tsx",
        "ShopGodowns": "features/shop/ShopGodownsScreen.tsx",
        "ShopLoyalty": "features/shop/ShopLoyaltyScreen.tsx",
        "ShopOrderDetail": "features/shop/ShopOrderDetailScreen.tsx",
        "ShopBooksParties": "features/shop/ShopBooksPartiesScreen.tsx",
        "ShopBooksExpense": "features/shop/ShopBooksExpenseScreen.tsx",
        "ShopBooksCash": "features/shop/ShopBooksCashScreen.tsx",
        "ShopBooksCheques": "features/shop/ShopBooksChequesScreen.tsx",
        "ShopBooksLoans": "features/shop/ShopBooksLoansScreen.tsx",
        "ShopBooksDocuments": "features/shop/ShopBooksDocumentsScreen.tsx",
        "ShopBooksCompliance": "features/shop/ShopBooksComplianceScreen.tsx",
        "ShopBooksSale": "features/shop/ShopBooksSaleScreen.tsx",
        "GrowWhatsApp": "features/grow/WhatsAppScreen.tsx",
        "GrowGoogleProfile": "features/grow/GoogleProfileScreen.tsx",
        "GrowAds": "features/grow/GrowAdsScreen.tsx",
        "GrowReferral": "features/grow/GrowReferralScreen.tsx",
        "PlatformAdminTenantDetail": "features/admin/PlatformAdminTenantDetailScreen.tsx",
        "SupportTicketDetail": "features/support/SupportTicketDetailScreen.tsx",
    }
    for (area, screen, route, mode, how), count in seen.items():
        screens.append([area, screen, route, mode, how, count, file_hint.get(route, "")])
    style_header(screens, len(sh_headers))
    autosize(screens, [14, 32, 28, 14, 40, 12, 52])
    for row in screens.iter_rows(min_row=2, max_row=screens.max_row, max_col=7):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # By area counts
    area_ws = wb.create_sheet("By area")
    area_ws.append(["Area", "Fields"])
    counts = {}
    for r in ROWS:
        counts[r[0]] = counts.get(r[0], 0) + 1
    for area, n in counts.items():
        area_ws.append([area, n])
    style_header(area_ws, 2)
    autosize(area_ws, [18, 12])
    chart = BarChart()
    chart.type = "col"
    chart.title = "Fields per area"
    chart.y_axis.title = "Fields"
    data = Reference(area_ws, min_col=2, min_row=1, max_row=area_ws.max_row)
    cats = Reference(area_ws, min_col=1, min_row=2, max_row=area_ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.style = 10
    chart.legend = None
    chart.width = 18
    chart.height = 8
    area_ws.add_chart(chart, "D2")

    fields.sheet_view.showGridLines = False
    screens.sheet_view.showGridLines = False

    wb.save(OUT)
    print(f"Wrote {OUT} ({len(ROWS)} fields, {len(seen)} screens)")


if __name__ == "__main__":
    main()
